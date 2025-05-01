from flask import Flask, request, send_file, jsonify
from google.cloud import vision
from io import BytesIO
import openpyxl
from openpyxl.utils import get_column_letter
import re

app = Flask(__name__)

def clean_number(val):
    val = re.sub(r'[^\d,\.\-]', '', val).replace(',', '.').replace(' ', '')
    return val if re.match(r'^\d+(\.\d+)?$', val) else ''

def ocr_google_vision(image_bytes: bytes) -> str:
    client = vision.ImageAnnotatorClient()
    image = vision.Image(content=image_bytes)

    # Ustaw język OCR na polski
    image_context = vision.ImageContext(language_hints=["pl"])

    # Użyj dokładniejszego trybu OCR
    response = client.document_text_detection(image=image, image_context=image_context)

    # Pobierz cały tekst z OCR
    text = response.full_text_annotation.text if response.full_text_annotation else ""

    return text

def parse_ocr_text_new_format(ocr_text):
    lines = [line.strip() for line in ocr_text.split("\n") if line.strip()]
    produkty = []

    # 1. Wydziel nazw produktów
    nazwy = []
    ceny_rabaty_start = 0
    for i, line in enumerate(lines):
        if re.match(r"\d+\s*\*\s*\d{1,3},\d{2}", line) or re.match(r"^-?\d{1,3},\d{2}$", line):
            ceny_rabaty_start = i
            break
        if not any(word in line for word in ["PTU", "Kwota", "Suma", "Razem", "RAZEM"]):
            nazwy.append(line)

    # 2. Parsuj linie z cenami i rabatami (pozycje produktów)
    dane = lines[ceny_rabaty_start:]
    i = 0
    while i < len(dane):
        if re.match(r"\d+\s*\*\s*\d{1,3},\d{2}", dane[i]) or re.match(r".*\d{1,3},\d{2}.*\d{1,3},\d{2}.*", dane[i]):
            produkt = {
                "nazwa": nazwy[len(produkty)] if len(produkty) < len(nazwy) else f"Produkt {len(produkty)+1}",
                "ilosc": 1,
                "cena_jedn": 0,
                "cena_laczna": 0,
                "rabat": 0,
                "wlasciciel": "wspolne"
            }
            # liczby z tej linii
            liczby = re.findall(r"\d{1,3},\d{2}", dane[i])
            if len(liczby) >= 2:
                produkt["cena_jedn"] = float(liczby[-2].replace(",", "."))
                produkt["cena_laczna"] = float(liczby[-1].replace(",", "."))

            # poszukaj rabatów (mogą być kilka pod rząd)
            j = i + 1
            rabaty = 0
            while j < len(dane) and re.match(r"^-?\d{1,3},\d{2}$", dane[j]):
                rabaty += float(dane[j].replace(",", "."))
                j += 1

            produkt["rabat"] = -abs(rabaty)
            produkty.append(produkt)
            i = j
        else:
            i += 1

    return produkty

@app.route("/upload", methods=["POST"])
def upload():
    if 'file' not in request.files:
        return jsonify({"error": "Brak pliku"}), 400

    file = request.files['file']
    image_bytes = file.read()
    ocr_text = ocr_google_vision(image_bytes)
    return ocr_text, 200, {"Content-Type": "text/plain; charset=utf-8"}
    # produkty = parse_ocr_text_new_format(ocr_text)

    # wb = openpyxl.Workbook()
    # ws = wb.active
    # ws.title = "Paragon"

    # headers = ["Nazwa", "Ilosc", "Cena jednostkowa", "Cena laczna", "Rabat", "Do zaplaty", "Wlasciciel"]
    # ws.append(headers)

    # for prod in produkty:
    #     do_zaplaty = prod['cena_laczna'] + prod['rabat']
    #     ws.append([
    #         prod['nazwa'], prod['ilosc'], prod['cena_jedn'], prod['cena_laczna'],
    #         prod['rabat'], do_zaplaty, prod['wlasciciel']
    #     ])

    # for col in range(1, len(headers) + 1):
    #     ws.column_dimensions[get_column_letter(col)].width = 18
    # ws.freeze_panes = "A2"

    # last_row = len(produkty) + 1
    # ws[f"E{last_row + 2}"].value = "Suma moje:"
    # ws[f"F{last_row + 2}"].value = f"=SUMIF(G2:G{last_row},\"ja\",F2:F{last_row})"
    # ws[f"E{last_row + 3}"].value = "Suma ona:"
    # ws[f"F{last_row + 3}"].value = f"=SUMIF(G2:G{last_row},\"ona\",F2:F{last_row})"
    # ws[f"E{last_row + 4}"].value = "Suma wspolne:"
    # ws[f"F{last_row + 4}"].value = f"=SUMIF(G2:G{last_row},\"wspolne\",F2:F{last_row})"
    # ws[f"E{last_row + 6}"].value = "Ona ma mi oddac:"
    # ws[f"F{last_row + 6}"].value = (
    #     f"=SUMIF(G2:G{last_row},\"ona\",F2:F{last_row}) + SUMIF(G2:G{last_row},\"wspolne\",F2:F{last_row})/2"
    # )

    # output = BytesIO()
    # wb.save(output)
    # output.seek(0)
    # return send_file(
    #     output,
    #     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    #     as_attachment=True,
    #     download_name="paragon.xlsx"
    # )

if __name__ == "__main__":
    app.run(debug=True)
